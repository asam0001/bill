import os
import tempfile
import urllib.parse
import qrcode
from datetime import datetime

# ReportLab imports
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# openpyxl imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from utils.settings_service import get_all_settings
from utils.sales_service import get_bill_by_number, get_bill_items

def ensure_directory_exists(path):
    directory = os.path.dirname(path)
    if directory and not os.path.exists(directory):
        os.makedirs(directory)

def generate_invoice_pdf(bill_number, output_path=None):
    """
    Generates a professional PDF invoice for the given bill number.
    """
    bill = get_bill_by_number(bill_number)
    if not bill:
        raise ValueError(f"Bill #{bill_number} not found.")
        
    items = get_bill_items(bill_number)
    settings = get_all_settings()
    
    # Setup Output Path
    if output_path is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        output_path = os.path.join(base_dir, "invoices", f"invoice_{bill_number}.pdf")
        
    ensure_directory_exists(output_path)
    
    # Document settings
    doc = SimpleDocTemplate(output_path, pagesize=letter,
                            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    
    # Style Sheet
    styles = getSampleStyleSheet()
    
    # Custom styles
    style_shop_name = ParagraphStyle(
        'ShopName',
        parent=styles['Heading1'],
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#2C3E50'),
        alignment=0 # Left
    )
    style_header_details = ParagraphStyle(
        'HeaderDetails',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#7F8C8D')
    )
    style_bill_title = ParagraphStyle(
        'BillTitle',
        parent=styles['Heading2'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#16A085'),
        alignment=2 # Right
    )
    style_bill_meta = ParagraphStyle(
        'BillMeta',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        alignment=2 # Right
    )
    style_table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontSize=10,
        leading=12,
        textColor=colors.white,
        fontName='Helvetica-Bold'
    )
    style_table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontSize=9,
        leading=12
    )
    
    # Header Grid (Shop Info left, Invoice Details right)
    shop_info_text = f"""
    <b>{settings.get('shop_name', 'Medical Shop')}</b><br/>
    {settings.get('address', '123 Pharmacy St')}<br/>
    Phone: {settings.get('phone', '9999999999')}<br/>
    GSTIN: {settings.get('gst_number', 'GSTIN123')}
    """
    
    invoice_meta_text = f"""
    INVOICE / BILL<br/>
    <b>Bill No:</b> {bill['bill_number']}<br/>
    <b>Date:</b> {bill['date']}<br/>
    <b>Payment Mode:</b> {bill['payment_mode']}
    """
    
    header_data = [
        [Paragraph(shop_info_text, style_shop_name if i == 0 else style_header_details), 
         Paragraph(invoice_meta_text if i == 0 else "", style_bill_title if i == 0 else style_bill_meta)] 
        for i in range(2)
    ]
    # Restructure headers
    header_data = [
        [Paragraph(f"<b>{settings.get('shop_name', 'Medical Shop')}</b>", style_shop_name), Paragraph("INVOICE", style_bill_title)],
        [Paragraph(f"{settings.get('address', '123 Pharmacy St')}<br/>Phone: {settings.get('phone', '9999999999')}<br/>GSTIN: {settings.get('gst_number', 'GSTIN123')}", style_header_details),
         Paragraph(f"<b>Bill No:</b> {bill['bill_number']}<br/><b>Date:</b> {bill['date']}<br/><b>Payment:</b> {bill['payment_mode']}", style_bill_meta)]
    ]
    
    header_table = Table(header_data, colWidths=[320, 220])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 15))
    
    # Customer Details Section
    cust_name = bill['customer_name'] if bill['customer_name'] else "Walk-in Customer"
    cust_phone = bill['customer_phone'] if bill['customer_phone'] else "N/A"
    customer_text = f"<b>Billed To:</b> {cust_name} | <b>Phone:</b> {cust_phone}"
    story.append(Paragraph(customer_text, styles['Normal']))
    story.append(Spacer(1, 10))
    
    # Items Table Header
    table_data = [[
        Paragraph("S.No", style_table_header),
        Paragraph("Medicine / Batch", style_table_header),
        Paragraph("Qty (Tab)", style_table_header),
        Paragraph("Price/Tab", style_table_header),
        Paragraph("Disc %", style_table_header),
        Paragraph("Amount (₹)", style_table_header)
    ]]
    
    for idx, item in enumerate(items, 1):
        med_name_batch = f"<b>{item['medicine_name']}</b><br/>Batch: {item['batch']}"
        table_data.append([
            Paragraph(str(idx), style_table_cell),
            Paragraph(med_name_batch, style_table_cell),
            Paragraph(str(item['quantity']), style_table_cell),
            Paragraph(f"₹{item['price_per_tablet']:.2f}", style_table_cell),
            Paragraph(f"{item['discount_percent']}%", style_table_cell),
            Paragraph(f"₹{item['subtotal']:.2f}", style_table_cell)
        ])
        
    items_table = Table(table_data, colWidths=[40, 240, 60, 70, 50, 80])
    items_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#BDC3C7')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9F9')])
    ]))
    story.append(items_table)
    story.append(Spacer(1, 15))
    
    # Total calculation and UPI QR Code layout
    total_text = f"""
    <b>Cart Total:</b> ₹{bill['total_amount']:.2f}<br/>
    <b>Discount Amount ({bill['discount']}%):</b> ₹{bill['discount_amount']:.2f}<br/>
    <font size="12" color="#16A085"><b>Grand Total:</b> ₹{bill['grand_total']:.2f}</font>
    """
    
    qr_image = None
    upi_id = settings.get('upi_id', '').strip()
    
    # Generate UPI QR code if UPI ID exists
    if upi_id:
        shop_name_esc = urllib.parse.quote(settings.get('shop_name', 'Medical Shop'))
        upi_url = f"upi://pay?pa={upi_id}&pn={shop_name_esc}&am={bill['grand_total']:.2f}&cu=INR"
        
        qr = qrcode.QRCode(version=1, box_size=5, border=1)
        qr.add_data(upi_url)
        qr.make(fit=True)
        
        # Create temp file
        temp_dir = tempfile.gettempdir()
        qr_file_path = os.path.join(temp_dir, f"upi_qr_{bill_number}.png")
        
        qr_img = qr.make_image(fill_color="black", back_color="white")
        qr_img.save(qr_file_path)
        
        # Load as ReportLab Flowable Image
        qr_image = Image(qr_file_path, width=80, height=80)
        
    summary_data = []
    if qr_image:
        summary_data.append([
            Paragraph("Scan to Pay using UPI:", style_header_details),
            Paragraph(total_text, style_bill_meta)
        ])
        summary_data.append([
            qr_image,
            ""
        ])
        summary_table = Table(summary_data, colWidths=[200, 340])
        summary_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('SPAN', (1, 0), (1, 1)),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ]))
    else:
        summary_data.append(["", Paragraph(total_text, style_bill_meta)])
        summary_table = Table(summary_data, colWidths=[300, 240])
        summary_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        
    story.append(KeepTogether(summary_table))
    story.append(Spacer(1, 20))
    
    # Footer
    footer_style = ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=9,
        alignment=1, # Center
        textColor=colors.HexColor('#95A5A6')
    )
    story.append(Paragraph("Thank you for visiting! Wish you good health.", footer_style))
    
    # Build document
    doc.build(story)
    
    # Delete temporary QR code image if created
    if qr_image and os.path.exists(qr_file_path):
        try:
            os.remove(qr_file_path)
        except OSError:
            pass
            
    return output_path

def generate_report_excel(report_type, sales_data, top_sold, top_profitable, output_path=None):
    """
    Exports a comprehensive report to Excel.
    """
    if output_path is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        output_path = os.path.join(base_dir, "reports_export", f"report_{report_type}_{datetime.now().strftime('%Y%m%d')}.xlsx")
        
    ensure_directory_exists(output_path)
    
    wb = openpyxl.Workbook()
    
    # --- Sheet 1: Sales Summary ---
    ws_sales = wb.active
    ws_sales.title = "Sales Transactions"
    ws_sales.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws_sales.merge_cells("A1:J1")
    title_cell = ws_sales["A1"]
    title_cell.value = f"Sales Report - {report_type.capitalize()}"
    title_cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sales.row_dimensions[1].height = 40
    
    # Meta Info
    ws_sales["A3"] = "Generated Date:"
    ws_sales["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_sales["A3"].font = Font(bold=True)
    
    # Table Headers
    headers = [
        "Bill No", "Date", "Customer Name", "Customer Phone", 
        "Payment Mode", "Gross Total (₹)", "Discount %", 
        "Discount Amount (₹)", "Grand Total (₹)", "Profit (₹)"
    ]
    
    header_row = 5
    for col_idx, header in enumerate(headers, 1):
        cell = ws_sales.cell(row=header_row, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="16A085", end_color="16A085", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    ws_sales.row_dimensions[header_row].height = 25
    
    # Data Rows
    current_row = 6
    total_sales_sum = 0
    total_profit_sum = 0
    
    border_side = Side(style='thin', color='BDC3C7')
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    for bill in sales_data:
        ws_sales.cell(row=current_row, column=1, value=bill['bill_number']).alignment = Alignment(horizontal="center")
        ws_sales.cell(row=current_row, column=2, value=bill['date']).alignment = Alignment(horizontal="center")
        ws_sales.cell(row=current_row, column=3, value=bill['customer_name'] or "Walk-in")
        ws_sales.cell(row=current_row, column=4, value=bill['customer_phone'] or "N/A").alignment = Alignment(horizontal="center")
        ws_sales.cell(row=current_row, column=5, value=bill['payment_mode']).alignment = Alignment(horizontal="center")
        
        ws_sales.cell(row=current_row, column=6, value=bill['total_amount'])
        ws_sales.cell(row=current_row, column=7, value=bill['discount'])
        ws_sales.cell(row=current_row, column=8, value=bill['discount_amount'])
        ws_sales.cell(row=current_row, column=9, value=bill['grand_total'])
        ws_sales.cell(row=current_row, column=10, value=bill['total_profit'])
        
        total_sales_sum += bill['grand_total']
        total_profit_sum += bill['total_profit']
        
        # Apply formatting to numeric columns
        for c in [6, 7, 8, 9, 10]:
            ws_sales.cell(row=current_row, column=c).number_format = '#,##0.00'
            
        for c in range(1, 11):
            ws_sales.cell(row=current_row, column=c).border = thin_border
            
        current_row += 1
        
    # Total Summary Row
    ws_sales.cell(row=current_row, column=5, value="Total Summary:").font = Font(bold=True)
    ws_sales.cell(row=current_row, column=5).alignment = Alignment(horizontal="right")
    
    sales_total_cell = ws_sales.cell(row=current_row, column=9, value=total_sales_sum)
    sales_total_cell.font = Font(bold=True)
    sales_total_cell.number_format = '₹#,##0.00'
    
    profit_total_cell = ws_sales.cell(row=current_row, column=10, value=total_profit_sum)
    profit_total_cell.font = Font(bold=True)
    profit_total_cell.number_format = '₹#,##0.00'
    
    # Auto-adjust column widths
    for col in ws_sales.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_sales.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # --- Sheet 2: Top Selling & Profitable ---
    ws_top = wb.create_sheet("Top Medicines Analysis")
    ws_top.views.sheetView[0].showGridLines = True
    
    # Left Header: Top 10 Selling
    ws_top.merge_cells("A1:C1")
    ts_title = ws_top["A1"]
    ts_title.value = "Top 10 Selling Medicines"
    ts_title.font = Font(bold=True, color="FFFFFF", size=12)
    ts_title.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    ts_title.alignment = Alignment(horizontal="center")
    
    ws_top.cell(row=2, column=1, value="Medicine Name").font = Font(bold=True)
    ws_top.cell(row=2, column=2, value="Batch").font = Font(bold=True)
    ws_top.cell(row=2, column=3, value="Qty Sold (Tablets)").font = Font(bold=True)
    
    for r_idx, med in enumerate(top_sold, 3):
        ws_top.cell(row=r_idx, column=1, value=med['medicine_name'])
        ws_top.cell(row=r_idx, column=2, value=med['batch'])
        ws_top.cell(row=r_idx, column=3, value=med['total_sold'])
        for c in range(1, 4):
            ws_top.cell(row=r_idx, column=c).border = thin_border
            
    # Right Header: Top 10 Profitable (column E, F, G)
    ws_top.merge_cells("E1:G1")
    tp_title = ws_top["E1"]
    tp_title.value = "Top 10 Profitable Medicines"
    tp_title.font = Font(bold=True, color="FFFFFF", size=12)
    tp_title.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    tp_title.alignment = Alignment(horizontal="center")
    
    ws_top.cell(row=2, column=5, value="Medicine Name").font = Font(bold=True)
    ws_top.cell(row=2, column=6, value="Batch").font = Font(bold=True)
    ws_top.cell(row=2, column=7, value="Profit Generated (₹)").font = Font(bold=True)
    
    for r_idx, med in enumerate(top_profitable, 3):
        ws_top.cell(row=r_idx, column=5, value=med['medicine_name'])
        ws_top.cell(row=r_idx, column=6, value=med['batch'])
        p_cell = ws_top.cell(row=r_idx, column=7, value=med['total_profit'])
        p_cell.number_format = '₹#,##0.00'
        for c in [5, 6, 7]:
            ws_top.cell(row=r_idx, column=c).border = thin_border
            
    # Set columns widths
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws_top.column_dimensions[col].width = 20
        
    wb.save(output_path)
    return output_path

def generate_report_pdf(report_type, sales_data, top_sold, top_profitable, output_path=None):
    """
    Generates a professional sales PDF report.
    """
    if output_path is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        output_path = os.path.join(base_dir, "reports_export", f"report_{report_type}_{datetime.now().strftime('%Y%m%d')}.pdf")
        
    ensure_directory_exists(output_path)
    
    doc = SimpleDocTemplate(output_path, pagesize=letter,
                            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    styles = getSampleStyleSheet()
    
    # Custom Styles
    style_title = ParagraphStyle('RepTitle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor('#2C3E50'))
    style_subtitle = ParagraphStyle('RepSubTitle', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#7F8C8D'))
    style_section_h = ParagraphStyle('RepSecH', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor('#16A085'))
    style_table_h = ParagraphStyle('RepTableH', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.white, fontName='Helvetica-Bold')
    style_table_c = ParagraphStyle('RepTableC', parent=styles['Normal'], fontSize=8, leading=10)
    
    # Header
    story.append(Paragraph(f"Sales Report - {report_type.capitalize()}", style_title))
    story.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", style_subtitle))
    story.append(Spacer(1, 15))
    
    # Totals Summary Box
    tot_sales = sum(b['grand_total'] for b in sales_data)
    tot_profit = sum(b['total_profit'] for b in sales_data)
    
    summary_data = [
        [Paragraph(f"<b>Total Sales Volume:</b> ₹{tot_sales:.2f}", style_table_c), 
         Paragraph(f"<b>Total Margin/Profit:</b> ₹{tot_profit:.2f}", style_table_c)]
    ]
    summary_table = Table(summary_data, colWidths=[270, 270])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EAEDED')),
        ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#BDC3C7')),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 15))
    
    # Transactions Section
    story.append(Paragraph("Sales History Transactions", style_section_h))
    story.append(Spacer(1, 5))
    
    tx_headers = [
        Paragraph("Bill No", style_table_h),
        Paragraph("Date", style_table_h),
        Paragraph("Customer", style_table_h),
        Paragraph("Mode", style_table_h),
        Paragraph("Total (₹)", style_table_h),
        Paragraph("Disc %", style_table_h),
        Paragraph("Grand (₹)", style_table_h),
        Paragraph("Profit (₹)", style_table_h)
    ]
    tx_data = [tx_headers]
    for b in sales_data:
        tx_data.append([
            Paragraph(str(b['bill_number']), style_table_c),
            Paragraph(b['date'].split()[0], style_table_c),
            Paragraph(b['customer_name'] or "Walk-in", style_table_c),
            Paragraph(b['payment_mode'], style_table_c),
            Paragraph(f"₹{b['total_amount']:.2f}", style_table_c),
            Paragraph(f"{b['discount']}%", style_table_c),
            Paragraph(f"₹{b['grand_total']:.2f}", style_table_c),
            Paragraph(f"₹{b['total_profit']:.2f}", style_table_c)
        ])
        
    tx_table = Table(tx_data, colWidths=[40, 70, 110, 50, 70, 45, 75, 80])
    tx_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#BDC3C7')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9F9')]),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(tx_table)
    story.append(Spacer(1, 20))
    
    # Top 10 Sections
    # Layout Top 10 lists side-by-side using a Table container
    left_top_flow = []
    left_top_flow.append(Paragraph("Top 10 Selling Medicines", style_section_h))
    left_top_flow.append(Spacer(1, 5))
    
    left_headers = [Paragraph("Medicine", style_table_h), Paragraph("Batch", style_table_h), Paragraph("Qty", style_table_h)]
    left_data = [left_headers]
    for m in top_sold:
        left_data.append([
            Paragraph(m['medicine_name'], style_table_c),
            Paragraph(m['batch'], style_table_c),
            Paragraph(str(m['total_sold']), style_table_c)
        ])
    left_table = Table(left_data, colWidths=[130, 70, 50])
    left_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#16A085')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BDC3C7')),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    left_top_flow.append(left_table)
    
    right_top_flow = []
    right_top_flow.append(Paragraph("Top 10 Profitable Medicines", style_section_h))
    right_top_flow.append(Spacer(1, 5))
    
    right_headers = [Paragraph("Medicine", style_table_h), Paragraph("Batch", style_table_h), Paragraph("Profit (₹)", style_table_h)]
    right_data = [right_headers]
    for m in top_profitable:
        right_data.append([
            Paragraph(m['medicine_name'], style_table_c),
            Paragraph(m['batch'], style_table_c),
            Paragraph(f"₹{m['total_profit']:.2f}", style_table_c)
        ])
    right_table = Table(right_data, colWidths=[120, 60, 70])
    right_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2C3E50')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BDC3C7')),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    right_top_flow.append(right_table)
    
    # Master layout for top 10 columns
    columns_data = [[left_top_flow, right_top_flow]]
    columns_table = Table(columns_data, colWidths=[270, 270])
    columns_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('RIGHTPADDING', (0,0), (0,0), 10),
        ('LEFTPADDING', (1,0), (1,0), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(KeepTogether(columns_table))
    
    doc.build(story)
    return output_path
